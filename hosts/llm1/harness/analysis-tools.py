"""Stateless, network-free analysis jobs. All untrusted parsing/code runs in bwrap."""
import base64
import json
import os
from pathlib import Path
import re
import signal
import subprocess
import tempfile
import threading
import time
import uuid
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

PYTHON = os.environ.get('ANALYSIS_PYTHON', '/run/current-system/sw/bin/python3')
BWRAP = os.environ.get('ANALYSIS_BWRAP', '/run/current-system/sw/bin/bwrap')
ARTIFACTS = Path(os.environ.get('ANALYSIS_ARTIFACTS', '/tmp/strata-artifacts'))
SLOTS = threading.BoundedSemaphore(2)
MAX_FILE = 8_000_000
TOOLS = [{"name": "python", "description": "Execute Python for calculations, data analysis and charts in an isolated, network-free, disposable workspace. pandas, numpy, matplotlib, pypdf, docx and openpyxl are installed. Supply input files as a mapping of simple filenames to text contents. Save output files in the current directory; returned download links expire after 24 hours. State is NOT retained between calls. Never claim execution succeeded when an error is returned.", "inputSchema": {"type": "object", "properties": {"code": {"type": "string"}, "files": {"type": "object", "additionalProperties": {"type": "string"}}}, "required": ["code"], "additionalProperties": False}}]


def filename(name):
    if not isinstance(name, str) or not re.fullmatch(r'[\w .-]{1,120}', name) or name in ('.', '..'):
        raise ValueError('Use a simple filename without directory separators')
    return name


def sandbox(directory, code):
    command = [os.environ['ANALYSIS_PRLIMIT'], '--cpu=25', f'--fsize={MAX_FILE}', '--nofile=128', '--core=0', BWRAP, '--unshare-all', '--die-with-parent', '--new-session', '--cap-drop', 'ALL',
               '--ro-bind', '/nix/store', '/nix/store', '--dev', '/dev',
               '--tmpfs', '/tmp', '--dir', '/home', '--bind', str(directory), '/work', '--chdir', '/work',
               '--clearenv', '--setenv', 'HOME', '/tmp', '--setenv', 'MPLCONFIGDIR', '/tmp/matplotlib',
               '--setenv', 'MPLBACKEND', 'Agg', '--setenv', 'OPENBLAS_NUM_THREADS', '1',
               PYTHON, '-c', code]
    with tempfile.TemporaryFile() as output:
        p = subprocess.Popen(command, stdout=output, stderr=subprocess.STDOUT, start_new_session=True)
        try:
            p.wait(timeout=30)
        except subprocess.TimeoutExpired:
            os.killpg(p.pid, signal.SIGKILL)
            p.wait()
            raise ValueError('Analysis exceeded the 30 second execution limit')
        output.seek(0)
        text = output.read(16001).decode('utf-8', 'replace')
        return p.returncode, text[:16000], len(text) > 16000


def cleanup():
    ARTIFACTS.mkdir(mode=0o700, parents=True, exist_ok=True)
    for file in ARTIFACTS.iterdir():
        if file.is_file() and file.stat().st_mtime < time.time() - 86400:
            file.unlink(missing_ok=True)
    if sum(p.stat().st_size for p in ARTIFACTS.iterdir() if p.is_file()) > 256_000_000:
        raise ValueError('Artifact storage is full; retry after older downloads expire')


def execute(args):
    code = args.get('code')
    files = args.get('files', {})
    if not isinstance(code, str) or not 1 <= len(code) <= 64000:
        raise ValueError('code must contain 1-64000 characters')
    if not isinstance(files, dict) or len(files) > 10:
        raise ValueError('At most ten input files are supported')
    if not SLOTS.acquire(blocking=False):
        raise ValueError('Analysis workers busy; retry shortly')
    try:
        cleanup()
        with tempfile.TemporaryDirectory(prefix='analysis-') as directory:
            root = Path(directory)
            for name, content in files.items():
                filename(name)
                if not isinstance(content, str) or len(content.encode()) > MAX_FILE:
                    raise ValueError('Input file must be text under 8 MB')
                (root / name).write_text(content)
            status, output, truncated = sandbox(root, code)
            artifacts = []
            for path in sorted(root.iterdir()):
                if path.is_symlink() or not path.is_file() or path.stat().st_size > MAX_FILE:
                    continue
                if len(artifacts) >= 10:
                    break
                name = filename(path.name)
                token = uuid.uuid4().hex
                target = ARTIFACTS / token
                target.write_bytes(path.read_bytes())
                artifacts.append({'name': name, 'url': '/workspace/artifacts/' + token, 'bytes': target.stat().st_size})
            return {'exit_code': status, 'stdout': output, 'truncated': truncated, 'artifacts': artifacts,
                    'note': 'Downloads expire after 24 hours. Each call starts with an empty workspace.'}
    finally:
        SLOTS.release()


EXTRACT = r'''
import json
from pathlib import Path
p=Path('input')
ext=Path('name').read_text().lower().rsplit('.',1)[-1]
if ext=='pdf':
    from pypdf import PdfReader
    reader=PdfReader(p)
    if len(reader.pages)>200: raise ValueError('PDF exceeds 200 pages')
    pages=[page.extract_text() or '' for page in reader.pages]
    if not any(t.strip() for t in pages): raise ValueError('No extractable text; scanned PDFs require OCR')
    text='\n\n'.join(f'Page {i+1}\n'+text for i,text in enumerate(pages))
elif ext=='docx':
    from docx import Document
    d=Document(p)
    text='\n'.join(x.text for x in d.paragraphs)+'\n'+'\n'.join('\t'.join(c.text for c in r.cells) for t in d.tables for r in t.rows)
elif ext=='xlsx':
    from openpyxl import load_workbook
    w=load_workbook(p.open('rb'), read_only=True, data_only=True)
    if len(w.sheetnames)>20: raise ValueError('Workbook exceeds 20 sheets')
    parts=[]
    for s in w.worksheets:
        if (s.max_row or 0)>10000 or (s.max_column or 0)>100: raise ValueError('Worksheet exceeds 10000 rows or 100 columns; split the data')
        parts.append('Sheet: '+s.title)
        for row in s.iter_rows(max_row=10000, max_col=100, values_only=True):
            parts.append('\t'.join('' if v is None else str(v) for v in row))
    text='\n'.join(parts)
else: raise ValueError('Supported document formats: PDF, DOCX, XLSX')
if len(text)>200000: raise ValueError('Extracted text exceeds 200000 characters; split the document')
Path('extracted.txt').write_text(text)
'''


def extract(args):
    name = filename(args.get('name'))
    data = base64.b64decode(args.get('data', ''), validate=True)
    if not data or len(data) > MAX_FILE:
        raise ValueError('Document must be between 1 byte and 8 MB')
    if not SLOTS.acquire(blocking=False):
        raise ValueError('Analysis workers busy; retry shortly')
    try:
        with tempfile.TemporaryDirectory(prefix='extract-') as directory:
            root = Path(directory)
            (root / 'input').write_bytes(data)
            (root / 'name').write_text(name)
            status, output, _ = sandbox(root, EXTRACT)
            if status:
                raise ValueError('Document extraction failed: ' + output[-1000:])
            return {'name': name, 'text': (root / 'extracted.txt').read_text()}
    finally:
        SLOTS.release()


class Handler(BaseHTTPRequestHandler):
    def log_message(self, *_):
        pass

    def reply(self, code, data):
        body = json.dumps(data).encode()
        self.send_response(code)
        self.send_header('Content-Type', 'application/json')
        self.send_header('Content-Length', str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self):
        token = self.path.removeprefix('/artifacts/')
        if self.headers.get('Origin') or not self.path.startswith('/artifacts/') or not re.fullmatch('[0-9a-f]{32}', token):
            self.send_error(404); return
        path = ARTIFACTS / token
        if not path.is_file() or path.stat().st_mtime < time.time()-86400:
            self.send_error(404); return
        body = path.read_bytes()
        self.send_response(200)
        self.send_header('Content-Type', 'application/octet-stream')
        self.send_header('Content-Length', str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def do_POST(self):
        self.connection.settimeout(45)
        if self.headers.get('Origin') or self.path not in ('/mcp', '/extract'):
            self.send_error(403); return
        try:
            size = int(self.headers.get('Content-Length', '0'))
            if not 0 < size <= 12_000_000:
                raise ValueError('Invalid request size')
            req = json.loads(self.rfile.read(size))
            if self.path == '/extract':
                self.reply(200, extract(req)); return
            method = req['method']
            if 'id' not in req:
                self.send_response(202); self.end_headers(); return
            if method == 'initialize':
                result = {'protocolVersion':'2025-06-18','capabilities':{'tools':{}},'serverInfo':{'name':'analysis','version':'1.0.0'}}
            elif method == 'tools/list':
                result = {'tools':TOOLS}
            elif method == 'ping':
                result = {}
            elif method == 'tools/call':
                try:
                    if req['params']['name'] != 'python': raise ValueError('Unknown tool')
                    value = execute(req['params'].get('arguments',{}))
                    result = {'isError':value['exit_code'] != 0,'content':[{'type':'text','text':json.dumps(value, ensure_ascii=False)}]}
                except Exception as e:
                    result = {'isError':True,'content':[{'type':'text','text':str(e)[:1000]}]}
            else:
                self.reply(200, {'jsonrpc':'2.0','id':req['id'],'error':{'code':-32601,'message':'Unknown method'}}); return
            self.reply(200, {'jsonrpc':'2.0','id':req['id'],'result':result})
        except (ValueError, KeyError, TypeError) as e:
            self.reply(400, {'error':str(e)[:1000]})


if __name__ == '__main__':
    ThreadingHTTPServer(('127.0.0.1',8092), Handler).serve_forever()
